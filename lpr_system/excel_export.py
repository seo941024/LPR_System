"""
excel_export.py — 기간별 기록 엑셀 내보내기
"""

import os
from datetime import datetime
from typing import Optional

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)

EXPORT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")


def _header_style():
    return {
        "font":      Font(bold=True, color="FFFFFF", size=11),
        "fill":      PatternFill("solid", fgColor="74AAD3"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border":    Border(
            bottom=Side(style="thin", color="CCCCCC"),
        ),
    }


def _apply(cell, styles: dict):
    for k, v in styles.items():
        setattr(cell, k, v)


def export_plate_logs(rows, start_date: str, end_date: str) -> str:
    """
    plate_logs → xlsx 파일 저장.
    반환: 저장된 파일 경로
    """
    os.makedirs(EXPORT_DIR, exist_ok=True)
    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(EXPORT_DIR, f"인식기록_{start_date}_{end_date}_{ts}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "인식 기록"

    headers = ["ID", "번호판", "색상", "타입", "신뢰도", "인식 시각", "카메라", "상태"]
    hs = _header_style()
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        _apply(cell, hs)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 22
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 10

    fill_white = PatternFill("solid", fgColor="FFFFFF")
    fill_gray  = PatternFill("solid", fgColor="F5F0EA")
    fill_red   = PatternFill("solid", fgColor="FDECEA")

    for i, r in enumerate(rows, 2):
        r = dict(r)
        status = "차단" if r.get("is_black") else ("등록" if r.get("is_white") else "미등록")
        row_data = [
            r.get("id"),
            r.get("plate_text"),
            r.get("plate_color"),
            "신형" if r.get("plate_type") == "NEW" else "구형",
            f"{float(r.get('confidence', 0)):.1%}",
            r.get("timestamp"),
            r.get("camera_id"),
            status,
        ]
        fill = fill_red if r.get("is_black") else (fill_gray if i % 2 == 0 else fill_white)
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.fill      = fill
            cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H1"

    wb.save(path)
    return path


def export_entry_exit(rows, start_date: str, end_date: str) -> str:
    """
    entry_exit_log → xlsx 파일 저장.
    반환: 저장된 파일 경로
    """
    os.makedirs(EXPORT_DIR, exist_ok=True)
    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(EXPORT_DIR, f"입출차기록_{start_date}_{end_date}_{ts}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "입출차 기록"

    headers = ["ID", "번호판", "이벤트", "카메라", "시각", "소유자", "차량 설명", "연락처"]
    hs = _header_style()
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        _apply(cell, hs)

    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 16

    fill_entry = PatternFill("solid", fgColor="EAF3DE")
    fill_exit  = PatternFill("solid", fgColor="FAEEDA")
    fill_alt   = PatternFill("solid", fgColor="F5F0EA")

    for i, r in enumerate(rows, 2):
        r = dict(r)
        is_entry = r.get("event_type") == "ENTRY"
        row_data = [
            r.get("id"),
            r.get("plate_text"),
            "입차" if is_entry else "출차",
            r.get("camera_id"),
            r.get("timestamp"),
            r.get("owner_name") or "-",
            r.get("vehicle_desc") or "-",
            r.get("phone") or "-",
        ]
        fill = fill_entry if is_entry else fill_exit
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.fill      = fill
            cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:H1"

    wb.save(path)
    return path
